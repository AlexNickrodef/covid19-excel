import datetime
from flask import flash, send_file
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference


def prepare(country, date_from, date_to) -> str:
    """ Подготовка API запроса """

    query = 'https://api.covid19tracking.narrativa.com/api/country/:country?date_from=:date_from&date_to=:date_to'
    query = query.replace(':country', country).replace(':date_from', str(date_from)).replace(':date_to', str(date_to))
    return query


def country_is_in_file(country_for_check) -> bool:
    """ Проверка наличия страны в файле countries.txt """

    with open("utils/countries.txt", "r") as countries_file:
        # считать все страны из файла
        countries = countries_file.readlines()

        if (country_for_check.lower() + "\n") not in countries:
            # страна не найдена
            flash("I have no data for this country")
            return False
        else:
            # страна найдена
            return True


def make_charts(ws, data, country):
    """ Формирование графиков """

    # Формируем графики:
    dates_count = len(list(data['dates'].keys()))
    date_from = list(data['dates'].keys())[0]
    date_to = list(data['dates'].keys())[-1]

    # Указываем что даты лежат в col1
    dates = Reference(ws, min_col=1, min_row=1, max_row=dates_count + 1)

    # График 1
    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=dates_count + 1)
    chart = LineChart()
    chart.title = "Confirmed cases in " + country + " from " + date_from + " to " + date_to
    chart.style = 15
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 1 в отчет
    ws.add_chart(chart, "K1")

    # График 2
    data = Reference(ws, min_col=4, max_col=5, min_row=1, max_row=dates_count + 1)
    chart = LineChart()
    chart.title = "Active cases in " + country + " from " + date_from + " to " + date_to
    chart.style = 14
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 2 в отчет
    ws.add_chart(chart, "K15")

    # График 3
    data = Reference(ws, min_col=6, max_col=7, min_row=1, max_row=dates_count + 1)
    chart = LineChart()
    chart.title = "Recovered cases " + country + " from " + date_from + " to " + date_to
    chart.style = 13
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 3 в отчет
    ws.add_chart(chart, "T1")

    # График 4
    data = Reference(ws, min_col=8, max_col=9, min_row=1, max_row=dates_count + 1)
    chart = LineChart()
    chart.title = "Death cases in " + country + " from " + date_from + " to " + date_to
    chart.style = 12
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 4 в отчет
    ws.add_chart(chart, "T15")


def make_report(data, country, include_charts) -> str:
    """ Создание финального отчета """

    # Создаем workbook excel и забираем активный worksheet
    wb = Workbook()
    ws = wb.active

    rows = [('Date:', 'today confirmed cases',
             'new confirmed cases:', 'today active cases:',
             'new active cases:', 'today recovered cases:',
             'new recovered cases:', 'today deaths:', 'new deaths:',)]

    for date in data['dates']:
        statistic = data['dates'][date]['countries'][country]
        rows.append((date, statistic['today_confirmed'], statistic['today_new_confirmed'],
                     statistic['today_open_cases'], statistic['today_new_open_cases'],
                     statistic['today_recovered'], statistic['today_new_recovered'],
                     statistic['today_deaths'], statistic['today_new_deaths']))

    for row in rows:
        ws.append(row)

    if include_charts:
        # Если пользователь запросил графики, добавить их в отчет
        make_charts(ws, data, country)

    timestamp = datetime.datetime.now().strftime("%d-%m-%Y %H.%M.%S")
    report = "reports/" + timestamp + " " + country + ".xlsx"

    # Сохраняем отчет
    wb.save(report)

    return report


def check_input_for_void(country, date_from, date_to) -> bool:
    """ Проверка данных из формы.
    Если какое-то поле не заполнено - вывести соответсвующее сообщение и вернуть False
    Если все поля заполнены - вернуть True """

    if not country:
        # Если страна не указана - перезагрузить страницу
        flash('You must specify the country.')
        return False

    if not (date_from and date_to):
        # Если даты не указаны - перезагрузить страницу
        flash('You must specify dates.')
        return False

    return True


def response_code_is_200(code) -> bool:
    """ Если сервер ответил не кодом 200 - вывести сообщение об ошибке """

    if code != 200:
        flash("Oops. Server responded with code:", code)
        return False
    else:
        return True


def check_dates_for_validity(date_from, date_to) -> bool:
    """ Проверка введенных дат на валидность """

    if date_to < date_from:
        # Дата "От" находится после даты "До"
        flash('Date "From" can\'t be past date "To".')
        return False

    elif date_to.strftime('%Y-%m-%d') > datetime.datetime.today().strftime('%Y-%m-%d'):
        # Дата "До" находится в будущем
        flash("I can't predict future yet.")
        return False

    else:
        # Ошибок не найдено
        return True


def convert_str_to_date(string) -> datetime.date:
    """ Конвертирование строки в дату """

    year, month, day = map(int, string.split('-'))
    date = datetime.date(year, month, day)

    return date
