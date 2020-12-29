import requests
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
import datetime
from openpyxl.chart.axis import DateAxis


# Подготовка запроса к источнику данных
def prepare(endpoint, country, date_from, date_to):
    date_from = date_from.strftime("%Y-%m-%d")
    date_to = date_to.strftime("%Y-%m-%d")
    return endpoint.replace(':country', country).replace(':date_from', date_from).replace(':date_to', date_to)


def make_report(data):

    # Создаем workbook excel и забираем активный worksheet
    wb = Workbook()
    ws = wb.active

    country = ''

    rows = [('Date:', 'today confirmed cases',
             'new confirmed cases:', 'today active cases:',
             'new active cases:', 'today recovered cases:',
             'new recovered cases:', 'today deaths:', 'new deaths:',)]

    for date in data['dates']:
        if data['dates'][date]['countries']:
            country = list(data['dates'][date]['countries'].keys())[0]
        elif country == '':
            print('Указаная страна не найдена')
            exit()

        statistic = data['dates'][date]['countries'][country]
        rows.append((date, statistic['today_confirmed'], statistic['today_new_confirmed'], statistic['today_open_cases'], statistic['today_new_open_cases'], statistic['today_recovered'], statistic['today_new_recovered'], statistic['today_deaths'], statistic['today_new_deaths']))

    for row in rows:
        ws.append(row)

    # Формируем графики:
    dates_count = len(list(data['dates'].keys()))
    date_from = list(data['dates'].keys())[0]
    date_to = list(data['dates'].keys())[-1]

    # Указываем что даты лежат в col1
    dates = Reference(ws, min_col=1, min_row=1, max_row=dates_count+1)

    # График 1
    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=dates_count+1)
    chart = LineChart()
    chart.title = "Confirmed cases in " + country + " from " + date_from + " to " + date_to
    chart.style = 15
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 1 отчет
    ws.add_chart(chart, "K1")

    # График 2
    data = Reference(ws, min_col=4, max_col=5, min_row=1, max_row=dates_count+1)
    chart = LineChart()
    chart.title = "Active cases in " + country + " from " + date_from + " to " + date_to
    chart.style = 14
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 2 отчет
    ws.add_chart(chart, "K15")

    # График 3
    data = Reference(ws, min_col=6, max_col=7, min_row=1, max_row=dates_count+1)
    chart = LineChart()
    chart.title = "Recovered cases " + country + " from " + date_from + " to " + date_to
    chart.style = 13
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 3 отчет
    ws.add_chart(chart, "T1")

    # График 4
    data = Reference(ws, min_col=8, max_col=9, min_row=1, max_row=dates_count+1)
    chart = LineChart()
    chart.title = "Death cases in " + country + " from " + date_from + " to " + date_to
    chart.style = 12
    chart.y_axis.title = "count"
    chart.y_axis.crossAx = 100
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(dates)

    # Добавляем График 4 отчет
    ws.add_chart(chart, "T15")

    timestamp = datetime.datetime.now().strftime("%d-%m-%Y %H.%M.%S")
    report = timestamp + " " + country + ".xlsx"

    # Сохраняем отчет
    wb.save(report)

    return report


if __name__ == '__main__':

    # ex.: https://api.covid19tracking.narrativa.com/api/country/tajikistan?date_from=2020-12-1&date_to=2020-12-10
    endpoint = 'https://api.covid19tracking.narrativa.com/api/country/:country?date_from=:date_from&date_to=:date_to'

    # тут будут наши данные
    data = []

    country = input("Enter country from counties.txt:")

    date_entry = input('Enter date (d.m.Y) starts:')
    day, month, year = map(int, date_entry.split('.'))
    date_from = datetime.date(year, month, day)

    date_entry = input('Enter date (d.m.Y) ends:')
    day, month, year = map(int, date_entry.split('.'))
    date_to = datetime.date(year, month, day)

    print('Processing..')

    try:
        # Подставляем данные в запрос
        endpoint = prepare(endpoint, country, date_from, date_to)

        # Пытаемся сделать запрос на API postman
        response = requests.get(endpoint)

        if response.status_code != 200:
            # Если код ответа не 200, что-то не так
            print("Resource respone by code: ", response.status_code)
            exit()

        # Если ок - обрабатываем json
        data = response.json()
    except Exception as err:
        print("Something wrong: ", err)
        exit()

    # Создаем отчет excel
    filename = make_report(data)

    print("Report is generated and saved by:", filename)